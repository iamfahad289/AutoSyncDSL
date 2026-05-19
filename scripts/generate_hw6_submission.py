#!/usr/bin/env python3
"""Generate the complete HW6 preliminary results submission package."""

from __future__ import annotations

import csv
import json
import os
import platform
import subprocess
import sys
from pathlib import Path
from textwrap import dedent

import matplotlib.pyplot as plt
import matplotlib.patches as patches
import numpy as np
from PIL import Image, ImageDraw, ImageFont

from autosyncdsl import Executor, SensorReading, SyncPlan
from baselines.manual_sync import ManualSyncBaseline, ManualSyncVariant2, SimpleSync
from loaders.kitti_loader import KITTILoader
from loaders.kitti_perturbation import KITTIPerturbationGenerator, PerturbedScenario


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "final_submission_hw6"
FIG = OUT / "figures"
TAB = OUT / "tables"
LOG = OUT / "logs"
ASSETS = OUT / "assets"
RESULTS = ROOT / "results" / "all_experiments.json"

TITLE = (
    "Homework 6: Preliminary Results\n"
    "AutoSyncDSL: An Embedded Python DSL for Multi-Sensor Timestamp Alignment "
    "and Batching in AV Perception Pipelines"
)
AUTHOR = "Muhammad Fahad"


def ensure_dirs() -> None:
    for path in [OUT, FIG, TAB, LOG, ASSETS]:
        path.mkdir(parents=True, exist_ok=True)
    cleanup_hidden_files()


def cleanup_hidden_files() -> None:
    for ds_store in OUT.rglob(".DS_Store"):
        ds_store.unlink()


def load_results() -> dict:
    with RESULTS.open() as f:
        return json.load(f)


def configure_plots() -> None:
    plt.rcParams.update(
        {
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
            "font.family": "Times New Roman",
            "font.size": 11,
            "axes.titlesize": 13,
            "axes.labelsize": 11,
            "xtick.labelsize": 10,
            "ytick.labelsize": 10,
            "legend.fontsize": 9,
            "text.color": "black",
            "axes.labelcolor": "black",
            "axes.edgecolor": "black",
            "xtick.color": "black",
            "ytick.color": "black",
        }
    )


def save_fig(fig, name: str) -> None:
    fig.savefig(FIG / name, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_box(ax, xy, width, height, title, body=None, fill="#f3f4f6", edge="#222222"):
    x, y = xy
    rect = patches.FancyBboxPatch(
        (x, y),
        width,
        height,
        boxstyle="round,pad=0.018,rounding_size=0.025",
        linewidth=1.3,
        edgecolor=edge,
        facecolor=fill,
    )
    ax.add_patch(rect)
    ax.text(x + width / 2, y + height - 0.07, title, ha="center", va="top", weight="bold")
    if body:
        lines = body if isinstance(body, list) else body.split("\n")
        start = y + height - 0.18
        for i, line in enumerate(lines):
            ax.text(x + 0.03, start - i * 0.055, line, ha="left", va="top", fontsize=9)


def arrow(ax, start, end, label=None):
    ax.annotate("", xy=end, xytext=start, arrowprops=dict(arrowstyle="->", lw=1.2, color="black"))
    if label:
        ax.text((start[0] + end[0]) / 2, (start[1] + end[1]) / 2 + 0.035, label, ha="center", fontsize=9)


def fig1_system_overview() -> None:
    fig, ax = plt.subplots(figsize=(11.2, 5.6))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.set_title("Figure 1. AutoSyncDSL System Overview", pad=10, weight="bold")

    draw_box(ax, (0.03, 0.56), 0.14, 0.30, "Sensor Streams", ["Camera", "LiDAR", "IMU"], "#f8fafc")
    draw_box(
        ax,
        (0.23, 0.56),
        0.17,
        0.30,
        "Embedded DSL",
        ['SyncPlan()', '.nearest(...)', ".drop_stale(...)", ".batch(...)"],
        "#e0f2fe",
    )
    draw_box(ax, (0.46, 0.56), 0.16, 0.30, "IR Lowering", ["source nodes", "match nodes", "filter node", "batch node"], "#e5e7eb")
    draw_box(ax, (0.68, 0.56), 0.14, 0.30, "Runtime", ["buffer", "match", "filter", "assemble"], "#dcfce7")
    draw_box(ax, (0.87, 0.56), 0.10, 0.30, "Output", ["sync", "batches"], "#fef3c7")

    arrow(ax, (0.17, 0.71), (0.23, 0.71), "specify")
    arrow(ax, (0.40, 0.71), (0.46, 0.71), "compile")
    arrow(ax, (0.62, 0.71), (0.68, 0.71), "execute")
    arrow(ax, (0.82, 0.71), (0.87, 0.71), "emit")

    draw_box(
        ax,
        (0.16, 0.14),
        0.68,
        0.24,
        "Implemented HW6 Path",
        [
            "DSL source builds a structured IR instead of executing immediately.",
            "Executor supports nearest-match alignment, stale filtering, IMU interpolation hooks, and batching.",
            "Evaluation scripts compare DSL behavior with handwritten baselines on synthetic and limited KITTI data.",
        ],
        "#ffffff",
    )
    save_fig(fig, "fig1_main_system_overview.png")


def fig2_lowering() -> None:
    fig, ax = plt.subplots(figsize=(11.2, 5.6))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.set_title("Figure 2. DSL-to-IR Lowering", pad=10, weight="bold")
    code = [
        "SyncPlan()",
        '  .camera("cam")',
        '  .lidar("lidar")',
        '  .imu("imu")',
        '  .nearest("cam", tolerance_ms=50)',
        "  .drop_stale(max_age_ms=100)",
        "  .batch(size=4)",
    ]
    draw_box(ax, (0.04, 0.18), 0.28, 0.64, "DSL Source", code, "#f8fafc")
    draw_box(ax, (0.40, 0.18), 0.26, 0.64, "IR Nodes", ["SensorSource", "NearestMatch", "StaleFilter", "Batch"], "#e5e7eb")
    draw_box(ax, (0.74, 0.18), 0.22, 0.64, "Executor Plan", ["topological node order", "configured tolerance", "batch size metadata"], "#dcfce7")
    arrow(ax, (0.32, 0.50), (0.40, 0.50), "compile()")
    arrow(ax, (0.66, 0.50), (0.74, 0.50), "run()")
    save_fig(fig, "fig2_dsl_to_ir_lowering.png")


def fig3_sync_semantics() -> None:
    fig, axes = plt.subplots(4, 1, figsize=(10.5, 8.5), sharex=True)
    cam = np.array([0, 33, 67, 100, 133])
    lidar = np.array([10, 50, 90, 120])
    colors = {"cam": "#1f77b4", "lidar": "#d97706", "imu": "#2ca02c", "bad": "#b91c1c"}
    titles = [
        "(a) Exact match: different sampling rates rarely share identical timestamps",
        "(b) Nearest match: LiDAR readings selected inside a 50 ms window",
        "(c) Stale rejection: readings older than max_age are rejected",
        "(d) IMU interpolation: high-rate samples support intermediate values",
    ]
    for ax, title in zip(axes, titles):
        ax.set_facecolor("white")
        ax.set_title(title, loc="left", fontsize=11)
        ax.grid(axis="x", color="#d1d5db", linewidth=0.6)
        ax.set_xlim(-5, 150)
        ax.tick_params(axis="y", length=0)

    axes[0].scatter(cam, [1] * len(cam), marker="o", color=colors["cam"], label="Camera")
    axes[0].scatter(lidar, [0] * len(lidar), marker="s", color=colors["lidar"], label="LiDAR")
    axes[0].set_yticks([0, 1], ["LiDAR", "Camera"])
    axes[0].legend(loc="upper right", frameon=False, ncol=2)

    axes[1].scatter(cam, [1] * len(cam), marker="o", color=colors["cam"], label="Camera")
    axes[1].scatter(lidar, [0] * len(lidar), marker="s", color=colors["lidar"], label="LiDAR")
    for c, l in zip(cam[:4], lidar[:4]):
        axes[1].plot([c, l], [1, 0], color="black", linestyle="--", linewidth=0.9)
    axes[1].set_yticks([0, 1], ["LiDAR", "Camera"])
    axes[1].text(122, 0.58, "within tolerance", fontsize=9)

    axes[2].scatter(cam, [1] * len(cam), marker="o", color=colors["cam"])
    axes[2].scatter([10, 50, 90], [0, 0, 0], marker="s", color=colors["lidar"])
    axes[2].scatter([10], [0], marker="x", color=colors["bad"], s=110, linewidths=2.4, label="stale")
    axes[2].axvline(133, color="black", linestyle=":", linewidth=1.0)
    axes[2].text(16, 0.20, "age > 100 ms", color=colors["bad"], fontsize=9)
    axes[2].set_yticks([0, 1], ["LiDAR", "Camera"])

    imu_ts = np.arange(0, 151, 10)
    imu_val = np.sin(imu_ts / 23.0)
    axes[3].plot(imu_ts, imu_val, color=colors["imu"], marker="o", linewidth=1.4, markersize=4, label="raw IMU")
    axes[3].axvline(65, color="black", linestyle="--", linewidth=1.0)
    axes[3].scatter([65], [np.sin(65 / 23.0)], color="black", zorder=5, label="interpolated value")
    axes[3].set_yticks([])
    axes[3].set_xlabel("Time (milliseconds)")
    axes[3].legend(loc="upper right", frameon=False)
    fig.suptitle("Figure 3. Synchronization Semantics", y=0.995, weight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.97))
    save_fig(fig, "fig3_sync_semantics.png")


def fig4_optimization() -> None:
    fig, ax = plt.subplots(figsize=(10.8, 5.4))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.set_title("Figure 4. Prototype Optimization Before and After", pad=10, weight="bold")
    draw_box(ax, (0.07, 0.18), 0.34, 0.64, "Before", ["NearestMatch pass", "StaleFilter pass", "Interpolation pass", "Batch assembly"], "#fee2e2")
    draw_box(ax, (0.59, 0.18), 0.34, 0.64, "After", ["candidate rule fusion", "buffer reuse metadata", "same user-facing semantics", "latency benefit still preliminary"], "#dcfce7")
    arrow(ax, (0.41, 0.50), (0.59, 0.50), "optimizer")
    ax.text(0.5, 0.08, "HW6 status: optimization passes exist, but measured latency improvement is not yet consistent.", ha="center", fontsize=10)
    save_fig(fig, "fig4_optimization_before_after.png")


def fig5_workflow() -> None:
    fig, ax = plt.subplots(figsize=(11.2, 5.3))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.set_title("Figure 5. Evaluation Workflow", pad=10, weight="bold")
    draw_box(ax, (0.04, 0.58), 0.20, 0.24, "Inputs", ["synthetic jitter", "dropout", "limited KITTI"], "#f8fafc")
    draw_box(ax, (0.33, 0.68), 0.20, 0.18, "AutoSyncDSL", ["DSL -> IR -> runtime"], "#e0f2fe")
    draw_box(ax, (0.33, 0.38), 0.20, 0.18, "Baselines", ["handwritten scripts"], "#fef3c7")
    draw_box(ax, (0.62, 0.53), 0.18, 0.24, "Metrics", ["latency", "group counts", "agreement"], "#dcfce7")
    draw_box(ax, (0.86, 0.53), 0.11, 0.24, "Outputs", ["JSON", "CSV", "figures"], "#e5e7eb")
    arrow(ax, (0.24, 0.70), (0.33, 0.77), "run")
    arrow(ax, (0.24, 0.62), (0.33, 0.47), "run")
    arrow(ax, (0.53, 0.77), (0.62, 0.67), "compare")
    arrow(ax, (0.53, 0.47), (0.62, 0.59), "compare")
    arrow(ax, (0.80, 0.65), (0.86, 0.65), "save")
    save_fig(fig, "fig5_eval_workflow.png")


def latency_mean(results: dict, scenario: str, method: str):
    for row in results["latency"].get(scenario, []):
        if row.get("approach") == method and "mean_ms" in row:
            return float(row["mean_ms"])
    return None


def fig6_latency(results: dict) -> None:
    scenarios = [("clean", "Clean synthetic"), ("jittery", "Jittery synthetic"), ("large", "Large synthetic"), ("kitti_real", "KITTI real")]
    methods = ["DSL", "DSL+Opt", "Baseline-v1", "Baseline-v2", "SimpleSync"]
    colors = {
        "DSL": "#1f77b4",
        "DSL+Opt": "#17becf",
        "Baseline-v1": "#d97706",
        "Baseline-v2": "#b91c1c",
        "SimpleSync": "#6b7280",
    }
    x = np.arange(len(scenarios))
    width = 0.15
    fig, ax = plt.subplots(figsize=(11.0, 5.8))
    for i, method in enumerate(methods):
        values = [latency_mean(results, key, method) for key, _ in scenarios]
        offsets = x + (i - 2) * width
        for ox, val in zip(offsets, values):
            if val is not None:
                ax.bar(ox, val, width, color=colors[method], label=method if ox == offsets[0] else None)
            else:
                ax.text(ox, 0.008, "n/a", ha="center", va="bottom", fontsize=7, rotation=90)
    ax.set_title("Figure 6. Preliminary Latency Comparison", weight="bold")
    ax.set_ylabel("Mean latency (milliseconds)")
    ax.set_xlabel("Scenario")
    ax.set_xticks(x)
    ax.set_xticklabels([label for _, label in scenarios])
    ax.grid(axis="y", color="#d1d5db", linewidth=0.6)
    ax.legend(frameon=False, ncol=5, loc="upper center", bbox_to_anchor=(0.5, 1.08))
    ax.text(
        0.01,
        0.96,
        "Lower is faster. Missing bars indicate methods not run for that scenario.",
        transform=ax.transAxes,
        ha="left",
        va="top",
        fontsize=9,
    )
    save_fig(fig, "fig6_latency_comparison.png")


def fig7_robustness(results: dict) -> None:
    jitter = results["robustness"]["jitter"]
    xs = [r["jitter_ms"] for r in jitter]
    dsl = [r["dsl_groups"] for r in jitter]
    base = [r["baseline_groups"] for r in jitter]
    fig, ax = plt.subplots(figsize=(9.5, 5.2))
    ax.plot(xs, dsl, marker="o", linewidth=2.0, color="#1f77b4", label="DSL")
    ax.plot(xs, base, marker="s", linewidth=1.6, linestyle="--", color="#d97706", label="Baseline")
    ax.set_title("Figure 7. Robustness Under Timestamp Jitter", weight="bold")
    ax.set_xlabel("Injected jitter (milliseconds)")
    ax.set_ylabel("Synchronized groups")
    ax.set_ylim(min(dsl + base) - 3, max(dsl + base) + 5)
    ax.grid(True, color="#d1d5db", linewidth=0.6)
    ax.legend(frameon=False)
    if dsl == base:
        ax.annotate(
            "Curves overlap exactly: DSL and baseline both produce 108 groups at every tested jitter level.",
            xy=(xs[len(xs) // 2], dsl[len(dsl) // 2]),
            xytext=(8, max(dsl) + 3),
            arrowprops=dict(arrowstyle="->", lw=1.0, color="black"),
            fontsize=9,
        )
    save_fig(fig, "fig7_robustness.png")


def count_lines(path: Path) -> int:
    return len(path.read_text().splitlines())


def fig8_complexity() -> None:
    dsl_parts = {
        "DSL API": count_lines(ROOT / "autosyncdsl" / "dsl.py"),
        "IR": count_lines(ROOT / "autosyncdsl" / "ir.py"),
        "Runtime": count_lines(ROOT / "autosyncdsl" / "executor.py"),
        "Optimization/support": count_lines(ROOT / "autosyncdsl" / "optimizations.py")
        + count_lines(ROOT / "autosyncdsl" / "utils.py")
        + count_lines(ROOT / "autosyncdsl" / "__init__.py"),
    }
    baseline_parts = {
        "Baseline implementations": count_lines(ROOT / "baselines" / "manual_sync.py"),
        "Package support": count_lines(ROOT / "baselines" / "__init__.py"),
    }
    labels = ["DSL+Runtime", "Handwritten Baselines"]
    colors = ["#1f77b4", "#6b7280", "#17becf", "#d97706"]
    fig, ax = plt.subplots(figsize=(8.4, 5.4))
    bottoms = [0, 0]
    all_keys = list(dsl_parts.keys())
    for i, key in enumerate(all_keys):
        val = dsl_parts[key]
        ax.bar(labels[0], val, bottom=bottoms[0], color=colors[i], label=key)
        bottoms[0] += val
    ax.bar(labels[1], baseline_parts["Baseline implementations"], bottom=0, color="#f59e0b", label="Baseline code")
    ax.bar(labels[1], baseline_parts["Package support"], bottom=baseline_parts["Baseline implementations"], color="#fde68a", label="Baseline package support")
    totals = [sum(dsl_parts.values()), sum(baseline_parts.values())]
    for label, total in zip(labels, totals):
        ax.text(label, total + 25, str(total), ha="center", fontsize=10, weight="bold")
    ax.set_title("Figure 8. Implementation Footprint", weight="bold")
    ax.set_ylabel("Physical lines of code")
    ax.grid(axis="y", color="#d1d5db", linewidth=0.6)
    ax.legend(frameon=False, bbox_to_anchor=(1.02, 1), loc="upper left")
    ax.text(0.02, 0.95, "Current DSL infrastructure is larger than the baseline module.", transform=ax.transAxes, fontsize=9, va="top")
    save_fig(fig, "fig8_code_complexity.png")


def generate_figures(results: dict) -> None:
    configure_plots()
    fig1_system_overview()
    fig2_lowering()
    fig3_sync_semantics()
    fig4_optimization()
    fig5_workflow()
    fig6_latency(results)
    fig7_robustness(results)
    fig8_complexity()


def run_for_group_count(method: str, cam, lidar, imu) -> int:
    if method == "DSL":
        plan = SyncPlan().camera("cam").lidar("lidar").imu("imu").nearest("lidar", tolerance_ms=50).drop_stale(max_age_ms=100).batch(size=4)
        batches = Executor(plan.compile()).run(cam, lidar, imu)
    elif method == "DSL+Opt":
        plan = SyncPlan().camera("cam").lidar("lidar").imu("imu").nearest("lidar", tolerance_ms=50).drop_stale(max_age_ms=100).batch(size=4)
        from autosyncdsl import Optimizer

        batches = Executor(Optimizer().optimize(plan.compile())).run(cam, lidar, imu)
    elif method == "Baseline-v1":
        batches = ManualSyncBaseline(match_type="nearest", tolerance_ms=50, max_age_ms=100, batch_size=4).synchronize(cam, lidar, imu)
    elif method == "Baseline-v2":
        batches = ManualSyncVariant2(tolerance_ms=50, max_age_ms=100, batch_size=4).synchronize(cam, lidar, imu)
    else:
        batches = SimpleSync(window_ms=100, batch_size=4).synchronize(cam, lidar, imu)
    return sum(len(batch) for batch in batches)


def scenario_streams(name: str):
    if name == "clean":
        return PerturbedScenario.clean_scenario()
    if name == "jittery":
        return PerturbedScenario.high_jitter_scenario()
    if name == "large":
        gen = KITTIPerturbationGenerator(duration_sec=30.0, camera_rate_hz=30, lidar_rate_hz=10, imu_rate_hz=100)
        return gen.generate(jitter_ms=10)
    if name == "kitti_real":
        loader = KITTILoader(str(ROOT / "data" / "raw" / "kitti"))
        drives = loader.load_drives()
        if drives:
            return loader.load_sequence(drives[0], max_frames=50)
    return None


def write_csv(path: Path, rows: list[dict], headers: list[str]) -> None:
    with path.open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def generate_tables(results: dict) -> tuple[list[dict], list[dict], list[dict]]:
    methods = ["DSL", "DSL+Opt", "Baseline-v1", "Baseline-v2", "SimpleSync"]
    scenarios = [("clean", "Clean synthetic"), ("jittery", "Jittery synthetic"), ("large", "Large synthetic"), ("kitti_real", "KITTI real")]
    metric_rows = []
    correctness_by_test = {r["test"]: r for r in results.get("correctness", []) if "test" in r}
    for key, label in scenarios:
        streams = scenario_streams(key)
        for method in methods:
            mean = latency_mean(results, key, method)
            if mean is None:
                continue
            groups = ""
            if key == "kitti_real":
                kitti = correctness_by_test.get("kitti_real", {})
                if method == "DSL" and kitti.get("status") == "ok":
                    groups = kitti["dsl"]["group_count"]
                elif method == "Baseline-v1" and kitti.get("status") == "ok":
                    groups = kitti["baseline"]["group_count"]
                else:
                    groups = "not recorded in correctness run"
            elif streams is not None:
                try:
                    groups = run_for_group_count(method, *streams)
                except Exception as exc:
                    groups = f"not recorded ({exc.__class__.__name__})"
            note = "latency mean from stored run; groups from HW6 table check"
            if key == "kitti_real":
                note = "limited 50-frame KITTI run; groups from correctness output when recorded"
            metric_rows.append(
                {
                    "Scenario": label,
                    "Method": method,
                    "Latency_ms": f"{mean:.4f}",
                    "Synchronized_Groups": groups,
                    "Notes": note,
                }
            )
    write_csv(
        TAB / "preliminary_metrics_table.csv",
        metric_rows,
        ["Scenario", "Method", "Latency_ms", "Synchronized_Groups", "Notes"],
    )

    literature_rows = [
        {
            "Paper/System": "ROS 2 message_filters ApproximateTimeSynchronizer",
            "Main Focus": "Approximate timestamp synchronization for message streams",
            "Relation to AutoSyncDSL": "Closest mature reference point for nearest-time synchronization policy",
            "Current Comparison Status": "Semantics are similar, but AutoSyncDSL is only a Python prototype and is not production middleware",
            "Why Different": "ROS integrates with runtime message queues; AutoSyncDSL currently emphasizes DSL, IR, and experimental execution",
        },
        {
            "Paper/System": "Geiger et al., The KITTI Vision Benchmark Suite, CVPR 2012",
            "Main Focus": "Autonomous-driving sensor dataset and benchmark",
            "Relation to AutoSyncDSL": "Provides the real-drive context and sensor-alignment motivation used in the project",
            "Current Comparison Status": "Only a limited local 50-frame KITTI sequence is processed at HW6",
            "Why Different": "KITTI is a benchmark dataset, not a synchronization DSL or middleware implementation",
        },
        {
            "Paper/System": "Lattner et al., MLIR, MAPS 2021",
            "Main Focus": "Compiler IR infrastructure and staged lowering",
            "Relation to AutoSyncDSL": "Motivates the DSL-to-IR separation at a smaller domain-specific scale",
            "Current Comparison Status": "Conceptually consistent, but AutoSyncDSL has a small custom IR and no general compiler backend",
            "Why Different": "MLIR is broad compiler infrastructure; AutoSyncDSL targets one synchronization domain",
        },
    ]
    write_csv(
        TAB / "literature_comparison_table.csv",
        literature_rows,
        ["Paper/System", "Main Focus", "Relation to AutoSyncDSL", "Current Comparison Status", "Why Different"],
    )

    plan_rows = [
        {
            "Task": "Investigate semantic discrepancies",
            "Deliverable": "Notes and fixes for stale filtering, high-dropout, and wide-tolerance mismatches",
            "Metric": "DSL/baseline agreement reported for each case",
            "Deadline": "May 5, 2026",
        },
        {
            "Task": "Finish optimization pass evaluation",
            "Deliverable": "Clean comparison of unoptimized and optimized IR execution",
            "Metric": "Mean and standard deviation over at least 20 runs per scenario",
            "Deadline": "May 8, 2026",
        },
        {
            "Task": "Broaden experiments",
            "Deliverable": "Additional synthetic scenarios and more KITTI frames if available",
            "Metric": "Latency, synchronized groups, offset statistics, and agreement",
            "Deadline": "May 11, 2026",
        },
        {
            "Task": "Finalize figures and tables",
            "Deliverable": "Final report-ready plots with captions and reproducibility notes",
            "Metric": "All figures generated from scripts and current JSON results",
            "Deadline": "May 14, 2026",
        },
        {
            "Task": "Write final report",
            "Deliverable": "Complete final paper and presentation-ready summary",
            "Metric": "Claims checked against logs, tables, and source code",
            "Deadline": "May 16, 2026",
        },
    ]
    write_csv(TAB / "final_report_plan_table.csv", plan_rows, ["Task", "Deliverable", "Metric", "Deadline"])
    return metric_rows, literature_rows, plan_rows


def table_to_md(rows: list[dict], headers: list[str]) -> str:
    widths = {h: max(len(h), *(len(str(r.get(h, ""))) for r in rows)) for h in headers}
    line1 = "| " + " | ".join(h.ljust(widths[h]) for h in headers) + " |"
    line2 = "| " + " | ".join("-" * widths[h] for h in headers) + " |"
    body = ["| " + " | ".join(str(r.get(h, "")).ljust(widths[h]) for h in headers) + " |" for r in rows]
    return "\n".join([line1, line2, *body])


def run_minimal_test() -> str:
    cam = [SensorReading(t, "cam", "camera", {"id": i}) for i, t in enumerate([0.0, 100.0, 200.0])]
    lidar = [SensorReading(t, "lidar", "lidar", {"id": i}) for i, t in enumerate([5.0, 105.0, 260.0])]
    imu = [SensorReading(t, "imu", "imu", {"id": i}) for i, t in enumerate([0.0, 100.0, 200.0])]
    plan = SyncPlan().camera("cam").lidar("lidar").imu("imu").nearest("cam", tolerance_ms=20).drop_stale(max_age_ms=50).batch(size=2)
    ir = plan.compile()
    batches = Executor(ir).run(cam, lidar, imu)
    groups = [g for batch in batches for g in batch]
    lines = [
        "Command: PYTHONPATH=. python <minimal AutoSyncDSL test>",
        "",
        "Test case:",
        "camera timestamps: [0.0, 100.0, 200.0]",
        "lidar timestamps:  [5.0, 105.0, 260.0]",
        "imu timestamps:    [0.0, 100.0, 200.0]",
        "policy: nearest(camera, tolerance_ms=20), drop_stale(max_age_ms=50), batch(size=2)",
        "",
        "Output:",
        f"synchronized batches: {len(batches)}",
        f"synchronized groups: {len(groups)}",
    ]
    for idx, group in enumerate(groups, start=1):
        lines.append(
            f"group {idx}: t={group.group_timestamp:.1f} ms, "
            f"camera={len(group.camera_readings)}, lidar={len(group.lidar_readings)}, imu={len(group.imu_readings)}"
        )
    lines.extend(["", "Compiled IR:", str(ir)])
    text = "\n".join(lines)
    (LOG / "example_test_case.txt").write_text(text)
    return text


def write_environment_info() -> None:
    packages = []
    for name in ["numpy", "matplotlib", "docx", "PIL"]:
        try:
            mod = __import__(name)
            packages.append(f"{name}: {getattr(mod, '__version__', 'available')}")
        except Exception as exc:
            packages.append(f"{name}: unavailable ({exc})")
    text = "\n".join(
        [
            "AutoSyncDSL HW6 environment information",
            f"Working directory: {ROOT}",
            "Primary command: PYTHONPATH=. python scripts/run_all_experiments.py",
            "Note: PYTHONPATH=. is required when invoking scripts by path from the repository root.",
            f"Python executable: {sys.executable}",
            f"Python version: {sys.version.split()[0]}",
            f"Platform: {platform.platform()}",
            f"Machine: {platform.machine()}",
            "Package versions:",
            *[f"  {p}" for p in packages],
            "Generated results path: results/all_experiments.json",
        ]
    )
    (LOG / "environment_info.txt").write_text(text)


def make_execution_evidence_image() -> None:
    log_path = LOG / "execution_log.txt"
    lines = log_path.read_text(errors="replace").splitlines()
    keep = [
        "$ PYTHONPATH=. python scripts/run_all_experiments.py",
        "============================================================",
        "AUTOSYNCDSL EVALUATION SUITE",
        "============================================================",
    ]
    patterns = [
        "Running: Exact Matching",
        "groups_dsl",
        "Generating clean scenario",
        "=== CLEAN SCENARIO ===",
        "DSL:",
        "  Mean:",
        "Baseline-v1:",
        "=== JITTER ROBUSTNESS ===",
        "2.0",
        "50.0",
        "Results saved to results/all_experiments.json",
        "EVALUATION COMPLETE",
    ]
    for line in lines:
        if any(p in line for p in patterns):
            keep.append(line)
        if len(keep) >= 34:
            break
    width, height = 1500, 900
    img = Image.new("RGB", (width, height), "#ffffff")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Menlo.ttc", 24)
        title_font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf", 34)
    except Exception:
        font = ImageFont.load_default()
        title_font = font
    draw.rectangle((0, 0, width, height), fill="#ffffff")
    draw.rectangle((40, 40, width - 40, height - 40), outline="#111111", width=2)
    draw.text((70, 65), "Execution Evidence: AutoSyncDSL HW6 Run", font=title_font, fill="#000000")
    y = 125
    for line in keep:
        if y > height - 70:
            break
        draw.text((70, y), line[:115], font=font, fill="#000000")
        y += 25
    img.save(FIG / "execution_evidence.png")


def selected_numbers(results: dict) -> dict:
    clean = {r["approach"]: r["mean_ms"] for r in results["latency"]["clean"]}
    kitti = {r["approach"]: r["mean_ms"] for r in results["latency"]["kitti_real"]}
    jitter = results["robustness"]["jitter"]
    correctness = {r["test"]: r for r in results["correctness"]}
    return {
        "clean_dsl": clean["DSL"],
        "clean_opt": clean["DSL+Opt"],
        "clean_base1": clean["Baseline-v1"],
        "clean_base2": clean["Baseline-v2"],
        "clean_simple": clean["SimpleSync"],
        "kitti_dsl": kitti["DSL"],
        "kitti_base1": kitti["Baseline-v1"],
        "jitter_min": min(r["jitter_ms"] for r in jitter),
        "jitter_max": max(r["jitter_ms"] for r in jitter),
        "jitter_groups": jitter[0]["dsl_groups"],
        "exact_groups": correctness["exact_matching"]["groups_dsl"],
        "stale_dsl": correctness["stale_filtering"]["dsl_kept"],
        "stale_base": correctness["stale_filtering"]["baseline_kept"],
        "kitti_groups": correctness["kitti_real"]["dsl"]["group_count"],
        "kitti_base_groups": correctness["kitti_real"]["baseline"]["group_count"],
    }


def build_markdown(results: dict, metric_rows: list[dict], literature_rows: list[dict], plan_rows: list[dict], minimal_test: str) -> str:
    n = selected_numbers(results)
    metric_preview = [r for r in metric_rows if r["Scenario"] in ("Clean synthetic", "KITTI real")]
    metric_compact = [
        {
            "Scenario": r["Scenario"],
            "Method": r["Method"],
            "Latency_ms": r["Latency_ms"],
            "Groups": r["Synchronized_Groups"],
        }
        for r in metric_preview
    ]
    literature_compact = [
        {
            "Reference": "ROS 2 message_filters",
            "Status": "Closest semantic comparison; AutoSyncDSL is less mature",
            "Difference": "Production middleware vs Python DSL prototype",
        },
        {
            "Reference": "KITTI Vision Benchmark Suite",
            "Status": "Dataset context; HW6 uses limited local sample",
            "Difference": "Benchmark dataset vs synchronization DSL",
        },
        {
            "Reference": "MLIR",
            "Status": "Conceptual IR/lowering reference",
            "Difference": "General compiler infrastructure vs small project IR",
        },
    ]
    plan_compact = [
        {"Task": r["Task"], "Metric": r["Metric"], "Deadline": r["Deadline"]}
        for r in plan_rows
    ]
    log_excerpt = "\n".join((LOG / "execution_log.txt").read_text(errors="replace").splitlines()[:18])
    md = f"""---
title: "Homework 6: Preliminary Results"
subtitle: "AutoSyncDSL: An Embedded Python DSL for Multi-Sensor Timestamp Alignment and Batching in AV Perception Pipelines"
author: "{AUTHOR}"
fontsize: 12pt
mainfont: "Times New Roman"
geometry: margin=1in
---

## 1. Progress Summary

Since HW5, I have moved AutoSyncDSL from a design-oriented prototype toward an executable preliminary evaluation pipeline: the embedded Python DSL is functional, plans lower into an inspectable IR, the runtime executor runs nearest-match synchronization, stale filtering, batching, and IMU interpolation hooks, and the repository now includes handwritten baseline scripts plus an evaluation pipeline for correctness, latency, and robustness experiments on synthetic perturbed streams and a limited KITTI real-drive sample.

## 2. Preliminary Results

The tested setup uses camera, LiDAR, and IMU timestamp streams generated by the repository's perturbed KITTI-style loader, plus a limited local KITTI sequence named in the execution log and capped at 50 frames. The main execution command used for this milestone was `PYTHONPATH=. python scripts/run_all_experiments.py`, which writes the aggregate JSON output used by the figures and tables. Figure 1 summarizes the implemented pipeline that is now exercised by the tests.

![](figures/fig1_main_system_overview.png)

*Figure 1. AutoSyncDSL system overview.*

Figure 3 shows the synchronization behavior used in the preliminary tests. Exact matching is included for completeness but is rarely realistic for heterogeneous sensors. The main current policy is nearest-neighbor matching within a tolerance window, followed by stale-frame rejection and batching. The worked example also shows where IMU interpolation fits into the semantics, although interpolation edge cases still need more testing before the final report.

![](figures/fig3_sync_semantics.png)

*Figure 3. Synchronization semantics.*

The preliminary latency results are mixed and should be interpreted as prototype timings rather than final performance claims. On the clean synthetic scenario, the DSL runtime averaged {n['clean_dsl']:.4f} ms, DSL+Opt averaged {n['clean_opt']:.4f} ms, Baseline-v1 averaged {n['clean_base1']:.4f} ms, Baseline-v2 averaged {n['clean_base2']:.4f} ms, and SimpleSync averaged {n['clean_simple']:.4f} ms. On the limited KITTI run, the DSL averaged {n['kitti_dsl']:.4f} ms compared with {n['kitti_base1']:.4f} ms for Baseline-v1. These numbers show that the current DSL implementation is not yet faster than the simpler baselines in the main comparisons; Python-level runtime overhead and incomplete optimization are likely contributing factors.

![](figures/fig6_latency_comparison.png)

*Figure 6. Preliminary latency comparison.*

The stronger preliminary result is robustness under timestamp jitter. Across jitter levels from {n['jitter_min']:.1f} ms to {n['jitter_max']:.1f} ms, the DSL and the primary handwritten baseline both produced {n['jitter_groups']} synchronized groups at every tested jitter level. This suggests that the current runtime preserves the intended nearest-match synchronization behavior under the tested jitter conditions. The broader robustness output is not uniformly solved yet: the execution log also shows disagreement at higher dropout rates and at the widest tolerance setting, so those cases are part of the final-report work rather than finished results.

![](figures/fig7_robustness.png)

*Figure 7. Robustness under timestamp jitter.*

Table 1 summarizes representative preliminary metrics from the generated CSV file. The synchronized-group counts are included to keep the timing results connected to actual produced outputs. The full table with notes is saved as `tables/preliminary_metrics_table.csv`.

Table 1. Representative preliminary metrics.

{table_to_md(metric_compact, ['Scenario', 'Method', 'Latency_ms', 'Groups'])}

Figure 8 is included only as an implementation-footprint figure. It does not claim that the current DSL reduces total code size. The opposite is true at this stage: the DSL, IR, runtime, and optimizer introduce infrastructure overhead. The value being tested is whether explicit semantics, modularity, and extensibility justify that overhead as the project matures.

![](figures/fig8_code_complexity.png)

*Figure 8. Implementation footprint.*

## 3. Comparison to Literature

The closest practical comparison from the HW5 reference set is ROS 2 `message_filters`, especially approximate-time synchronization for timestamped streams. AutoSyncDSL is consistent with that style of nearest-time alignment at the semantic level, but the current implementation is not comparable to mature middleware in performance, integration, or runtime robustness. The early results are therefore worse than a production system would be expected to achieve, mainly because AutoSyncDSL is a Python prototype with DSL/IR/runtime overhead, incomplete optimization, and a smaller evaluation setup. The KITTI Vision Benchmark Suite by Geiger et al. provides the dataset context for AV sensor streams, but this HW6 package only processes a limited local sample rather than conducting a full KITTI-scale evaluation. MLIR is also relevant as a compiler-design reference: AutoSyncDSL follows the idea of separating front-end syntax from an intermediate representation, but it uses a small project-specific IR rather than a general compiler framework.

Table 2. Literature and system comparison. The full comparison CSV is saved under `tables/`.

{table_to_md(literature_compact, ['Reference', 'Status', 'Difference'])}

References used in this comparison are Open Robotics ROS 2 `message_filters` package documentation, Geiger et al.'s KITTI Vision Benchmark Suite paper from CVPR 2012, and Lattner et al.'s MLIR paper from MAPS 2021.

## 4. Critical Reflection

Several parts remain incomplete or blocked. First, optimization is not yet producing a reliable latency improvement; in the clean run, DSL+Opt was slightly slower than the unoptimized DSL. Second, stale filtering still needs review because the current correctness output reports DSL kept {n['stale_dsl']} groups while the baseline kept {n['stale_base']} groups in that test. Third, high-dropout and very-wide-tolerance robustness cases show disagreement and need semantic inspection rather than being treated as successful results. Fourth, real-data evaluation is still limited to a small KITTI sample, where the DSL produced {n['kitti_groups']} groups and the baseline produced {n['kitti_base_groups']} groups. The main unexpected challenge is that making synchronization semantics explicit through a DSL and IR adds runtime and implementation overhead before the optimizer is strong enough to recover it. If the project stopped today, the main contribution would be a functioning DSL-to-IR-to-runtime pipeline that can express and execute synchronization policies and produce preliminary evidence of semantic agreement under controlled jitter.

## 5. Plan to Final Report

- Finish stale-filtering and dropout/tolerance discrepancy analysis by May 5, 2026, with a short explanation or patch for each mismatch.
- Re-run latency experiments with at least 20 repetitions per scenario by May 8, 2026, reporting mean, standard deviation, min, and max for DSL, DSL+Opt, and baselines.
- Complete optimization evaluation by May 8, 2026, including whether rule fusion or buffer reuse changes execution time on clean, jittery, large, and KITTI scenarios.
- Expand correctness and robustness tests by May 11, 2026, including synchronized-group counts, offset statistics, and agreement status across jitter, dropout, and tolerance settings.
- Add broader real-data evidence by May 11, 2026, if additional KITTI frames can be processed locally.
- Finalize report figures, captions, and result tables by May 14, 2026, with every numerical claim traceable to logs or generated CSV files.
- Complete the final report draft by May 16, 2026, keeping the contribution focused on DSL expressiveness, IR structure, runtime behavior, and honest performance tradeoffs.

Table 3. Final-report plan. The full planning CSV is saved under `tables/`.

{table_to_md(plan_compact, ['Task', 'Metric', 'Deadline'])}

## 6. Appendix: Evidence of Execution

The execution log in `logs/execution_log.txt` was generated from a real run of `PYTHONPATH=. python scripts/run_all_experiments.py`. The plain `python scripts/run_all_experiments.py` command fails in this environment because the repository root is not on the import path when the script is invoked by file path; this is documented in `logs/environment_info.txt`.

```text
{log_excerpt}
```

The minimal test case in `logs/example_test_case.txt` constructs three timestamp streams, compiles a DSL plan, runs the executor, and produces one synchronized batch containing two groups. The third camera timestamp is not synchronized because the nearest LiDAR reading is outside the 20 ms tolerance.

```text
{minimal_test.split('Compiled IR:')[0].strip()}
```

At HW6, the following pieces run: DSL construction, IR lowering, nearest-match execution, stale filtering, batching, baseline scripts, synthetic evaluation, limited KITTI loading, CSV/table generation, and figure generation. The following pieces are still incomplete or under validation: reliable optimization speedup, interpolation edge-case evaluation, full real-dataset coverage, and semantic reconciliation for high-dropout and very-wide-tolerance cases.

![](figures/execution_evidence.png)

*Execution evidence rendered from the run log.*

"""
    (OUT / "hw6.md").write_text(md)
    return md


def add_table_docx(doc, rows: list[dict], headers: list[str]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, header in enumerate(headers):
            cells[i].text = str(row.get(header, ""))


def build_docx(results: dict, metric_rows: list[dict], literature_rows: list[dict], plan_rows: list[dict], minimal_test: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    n = selected_numbers(results)
    metric_preview = [r for r in metric_rows if r["Scenario"] in ("Clean synthetic", "KITTI real")]
    metric_compact = [
        {
            "Scenario": r["Scenario"],
            "Method": r["Method"],
            "Latency_ms": r["Latency_ms"],
            "Groups": r["Synchronized_Groups"],
        }
        for r in metric_preview
    ]
    literature_compact = [
        {
            "Reference": "ROS 2 message_filters",
            "Status": "Closest semantic comparison; AutoSyncDSL is less mature",
            "Difference": "Production middleware vs Python DSL prototype",
        },
        {
            "Reference": "KITTI Vision Benchmark Suite",
            "Status": "Dataset context; HW6 uses limited local sample",
            "Difference": "Benchmark dataset vs synchronization DSL",
        },
        {
            "Reference": "MLIR",
            "Status": "Conceptual IR/lowering reference",
            "Difference": "General compiler infrastructure vs small project IR",
        },
    ]
    plan_compact = [{"Task": r["Task"], "Metric": r["Metric"], "Deadline": r["Deadline"]} for r in plan_rows]
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(12)
    for style_name in ["Title", "Heading 1", "Heading 2"]:
        styles[style_name].font.name = "Times New Roman"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Homework 6: Preliminary Results")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("AutoSyncDSL: An Embedded Python DSL for Multi-Sensor Timestamp Alignment and Batching in AV Perception Pipelines")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Author: {AUTHOR}")

    doc.add_heading("1. Progress Summary", level=1)
    doc.add_paragraph(
        "Since HW5, I have moved AutoSyncDSL from a design-oriented prototype toward an executable preliminary evaluation pipeline: the embedded Python DSL is functional, plans lower into an inspectable IR, the runtime executor runs nearest-match synchronization, stale filtering, batching, and IMU interpolation hooks, and the repository now includes handwritten baseline scripts plus an evaluation pipeline for correctness, latency, and robustness experiments on synthetic perturbed streams and a limited KITTI real-drive sample."
    )

    doc.add_heading("2. Preliminary Results", level=1)
    doc.add_paragraph(
        "The tested setup uses camera, LiDAR, and IMU timestamp streams generated by the repository's perturbed KITTI-style loader, plus a limited local KITTI sequence capped at 50 frames."
    )
    doc.add_picture(str(FIG / "fig1_main_system_overview.png"), width=Inches(6.4))
    doc.add_paragraph("Figure 1. AutoSyncDSL system overview.")
    doc.add_picture(str(FIG / "fig3_sync_semantics.png"), width=Inches(6.2))
    doc.add_paragraph("Figure 3. Synchronization semantics.")
    doc.add_paragraph(
        f"On the clean synthetic scenario, the DSL runtime averaged {n['clean_dsl']:.4f} ms, DSL+Opt averaged {n['clean_opt']:.4f} ms, Baseline-v1 averaged {n['clean_base1']:.4f} ms, Baseline-v2 averaged {n['clean_base2']:.4f} ms, and SimpleSync averaged {n['clean_simple']:.4f} ms. These are preliminary Python timings and do not show a final speedup claim."
    )
    doc.add_picture(str(FIG / "fig6_latency_comparison.png"), width=Inches(6.4))
    doc.add_paragraph("Figure 6. Preliminary latency comparison.")
    doc.add_paragraph(
        f"Across jitter levels from {n['jitter_min']:.1f} ms to {n['jitter_max']:.1f} ms, the DSL and primary baseline both produced {n['jitter_groups']} synchronized groups at every tested jitter level."
    )
    doc.add_picture(str(FIG / "fig7_robustness.png"), width=Inches(6.2))
    doc.add_paragraph("Figure 7. Robustness under timestamp jitter.")
    doc.add_picture(str(FIG / "fig8_code_complexity.png"), width=Inches(5.8))
    doc.add_paragraph("Figure 8. Implementation footprint.")
    doc.add_paragraph("The full preliminary metrics CSV includes notes for each row.")
    add_table_docx(doc, metric_compact, ["Scenario", "Method", "Latency_ms", "Groups"])

    doc.add_heading("3. Comparison to Literature", level=1)
    doc.add_paragraph(
        "The closest practical comparison from the HW5 reference set is ROS 2 message_filters, especially approximate-time synchronization for timestamped streams. AutoSyncDSL is consistent with nearest-time alignment semantically, but it is currently worse than mature middleware in performance and integration because it is a Python prototype with DSL/IR/runtime overhead, incomplete optimization, and a smaller evaluation setup. KITTI provides the AV dataset context, while MLIR motivates the DSL-to-IR separation at a small scale."
    )
    doc.add_paragraph("The full literature comparison CSV includes focus, relationship, status, and difference columns.")
    add_table_docx(doc, literature_compact, ["Reference", "Status", "Difference"])

    doc.add_heading("4. Critical Reflection", level=1)
    doc.add_paragraph(
        f"Optimization is not yet producing a reliable latency improvement, stale filtering still needs review because the current output reports DSL kept {n['stale_dsl']} groups while the baseline kept {n['stale_base']} groups, and high-dropout plus very-wide-tolerance cases show disagreement. If the project stopped today, the main contribution would be a functioning DSL-to-IR-to-runtime pipeline for expressing and executing synchronization policies."
    )

    doc.add_heading("5. Plan to Final Report", level=1)
    for row in plan_rows:
        doc.add_paragraph(f"{row['Task']}: {row['Deliverable']} ({row['Metric']}; deadline {row['Deadline']}).", style="List Bullet")
    add_table_docx(doc, plan_compact, ["Task", "Metric", "Deadline"])

    doc.add_heading("6. Appendix: Evidence of Execution", level=1)
    doc.add_paragraph("Execution command: PYTHONPATH=. python scripts/run_all_experiments.py")
    doc.add_paragraph(minimal_test.split("Compiled IR:")[0].strip())
    doc.add_picture(str(FIG / "execution_evidence.png"), width=Inches(6.4))
    doc.add_paragraph(
        "Runs now: DSL construction, IR lowering, nearest-match execution, stale filtering, batching, baseline scripts, synthetic evaluation, limited KITTI loading, tables, and figures. Still incomplete: reliable optimization speedup, interpolation edge-case evaluation, broader real data, and discrepancy reconciliation."
    )
    doc.save(OUT / "hw6.docx")


def build_pdf() -> None:
    cmd = [
        "pandoc",
        "hw6.md",
        "-o",
        "hw6.pdf",
        "--pdf-engine=xelatex",
        "-V",
        "mainfont=Times New Roman",
        "-V",
        "fontsize=12pt",
        "-V",
        "geometry:margin=1in",
    ]
    subprocess.run(cmd, cwd=OUT, check=True)


def write_manifest() -> None:
    text = """# Figure Manifest

All figures in this folder were regenerated for HW6 with white backgrounds, black text, readable labels, and high-resolution PNG export.

| File | Purpose | Included in HW6 PDF |
| --- | --- | --- |
| fig1_main_system_overview.png | System overview and implemented pipeline | Yes |
| fig2_dsl_to_ir_lowering.png | DSL source to IR lowering schematic | No |
| fig3_sync_semantics.png | Worked synchronization example | Yes |
| fig4_optimization_before_after.png | Prototype optimization status | No |
| fig5_eval_workflow.png | Evaluation workflow | No |
| fig6_latency_comparison.png | Quantitative latency comparison in milliseconds | Yes |
| fig7_robustness.png | Jitter robustness and DSL/baseline agreement | Yes |
| fig8_code_complexity.png | Implementation footprint, not a simplicity claim | Yes |
| execution_evidence.png | Terminal-style image rendered from the real execution log | Appendix |
"""
    (ASSETS / "figure_manifest.md").write_text(text)


def write_readme(results: dict) -> None:
    n = selected_numbers(results)
    text = f"""# HW6 Submission Package

This folder contains the complete Homework 6 preliminary results submission package for AutoSyncDSL.

## Checklist

- PDF exists: `hw6.pdf`
- DOCX exists: `hw6.docx`
- Markdown source exists: `hw6.md`
- Formatting: Times New Roman, 12 pt body text, black text, white page, single-column academic style
- Required sections are present in order: Progress Summary, Preliminary Results, Comparison to Literature, Critical Reflection, Plan to Final Report, Appendix: Evidence of Execution
- Figures are cleaned and readable under `figures/`
- Execution evidence is included in `logs/execution_log.txt`, `logs/example_test_case.txt`, `logs/environment_info.txt`, and `figures/execution_evidence.png`
- Claims are consistent with actual outputs from `results/all_experiments.json`
- Latency is reported honestly: clean DSL {n['clean_dsl']:.4f} ms, clean Baseline-v1 {n['clean_base1']:.4f} ms, clean Baseline-v2 {n['clean_base2']:.4f} ms
- Robustness is reported honestly: DSL and baseline both produce {n['jitter_groups']} groups for all tested jitter levels from {n['jitter_min']:.1f} ms to {n['jitter_max']:.1f} ms
- Code complexity is framed as implementation footprint, not reduced code size

## Remaining Before Final Report

- Resolve stale-filtering and high-dropout/tolerance semantic discrepancies.
- Re-run latency with more repetitions and report variance.
- Determine whether optimization passes provide measurable benefits.
- Broaden real-data evaluation beyond the current limited KITTI sample if possible.
- Keep all final claims traceable to logs, tables, and source code.

## Additional Verification Note

The HW6 evidence run completed successfully with `PYTHONPATH=. python scripts/run_all_experiments.py`. A separate repository unit-test check was attempted after packaging: the active Conda interpreter does not have `pytest`, and `.venv/bin/python -m pytest` is currently blocked by a syntax error in `tests/test_dsl.py` at line 320. This does not change the HW6 execution log, but it should be fixed before the final report.
"""
    (OUT / "README_hw6_submission.md").write_text(text)


def main() -> None:
    os.chdir(ROOT)
    ensure_dirs()
    results = load_results()
    generate_figures(results)
    metric_rows, literature_rows, plan_rows = generate_tables(results)
    minimal_test = run_minimal_test()
    write_environment_info()
    make_execution_evidence_image()
    build_markdown(results, metric_rows, literature_rows, plan_rows, minimal_test)
    build_docx(results, metric_rows, literature_rows, plan_rows, minimal_test)
    build_pdf()
    write_manifest()
    write_readme(results)
    cleanup_hidden_files()
    print(f"Generated HW6 package in {OUT}")


if __name__ == "__main__":
    main()
